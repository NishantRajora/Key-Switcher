import sys
import json
from pathlib import Path

from openpyxl import load_workbook

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QApplication,
    QComboBox,
    QFormLayout,
    QLabel,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QVBoxLayout,
    QHBoxLayout,
    QDialog,
    QWidget,
    QInputDialog,
)


# ============================================================
# FILE LOCATIONS
# ============================================================

SETTINGS_FILE = Path(
    r"C:\Users\Nishant\.claude\settings.json"
)

KEY_FILE = Path(
    r"C:\My Space\Data\key.xlsx"
)


# ============================================================
# AVAILABLE MODELS
# ============================================================

MODELS = [
    "gemma4:31b",
    "gpt-oss:120b",
    "gpt-oss:20b",
    "nemotron-3-nano:30b",
    "nemotron-3-super",
    "nemotron-3-ultra",
]


# ============================================================
# MODEL SETTINGS
# ============================================================

MODEL_SETTINGS = [
    "ANTHROPIC_MODEL",
    "ANTHROPIC_SMALL_FAST_MODEL",
    "ANTHROPIC_DEFAULT_SONNET_MODEL",
    "ANTHROPIC_DEFAULT_OPUS_MODEL",
    "ANTHROPIC_DEFAULT_HAIKU_MODEL",
]


# ============================================================
# WHITE THEME STYLESHEET
# ============================================================

APP_STYLESHEET = """
QMainWindow, QDialog {
    background-color: #ffffff;
}

QWidget {
    font-family: "Segoe UI", "Helvetica Neue", Arial, sans-serif;
    color: #1a1a1a;
}

QLabel {
    color: #1a1a1a;
}

QLabel#TitleLabel {
    font-size: 21px;
    font-weight: 600;
    color: #111111;
    padding-bottom: 4px;
}

QLabel#StatusLabel {
    font-size: 13px;
    color: #6b7280;
    padding-top: 4px;
}

QLabel#StatusLabel[state="success"] {
    color: #15803d;
}

QLabel#StatusLabel[state="error"] {
    color: #b91c1c;
}

QFormLayout QLabel {
    font-size: 13px;
    font-weight: 500;
    color: #374151;
}

QComboBox {
    background-color: #ffffff;
    border: 1px solid #d1d5db;
    border-radius: 6px;
    padding: 6px 10px;
    font-size: 13px;
    selection-background-color: #2563eb;
}

QComboBox:hover {
    border: 1px solid #9ca3af;
}

QComboBox:focus {
    border: 1px solid #2563eb;
}

QComboBox::drop-down {
    border: none;
    width: 24px;
}

QComboBox QAbstractItemView {
    background-color: #ffffff;
    border: 1px solid #d1d5db;
    selection-background-color: #eff6ff;
    selection-color: #1a1a1a;
    outline: none;
}

QLineEdit {
    background-color: #ffffff;
    border: 1px solid #d1d5db;
    border-radius: 6px;
    padding: 6px 10px;
    font-size: 13px;
}

QLineEdit:focus {
    border: 1px solid #2563eb;
}

QPushButton {
    background-color: #f3f4f6;
    border: 1px solid #d1d5db;
    border-radius: 6px;
    font-size: 13px;
    padding: 6px 12px;
}

QPushButton:hover {
    background-color: #e5e7eb;
}

QPushButton:pressed {
    background-color: #d1d5db;
}

QPushButton#ApplyButton {
    background-color: #2563eb;
    border: none;
    color: #ffffff;
    font-size: 15px;
    font-weight: 600;
    border-radius: 8px;
}

QPushButton#ApplyButton:hover {
    background-color: #1d4ed8;
}

QPushButton#ApplyButton:pressed {
    background-color: #1e40af;
}

QPushButton#DotsButton {
    background-color: transparent;
    border: none;
    font-size: 20px;
    font-weight: bold;
    color: #4b5563;
    border-radius: 18px;
}

QPushButton#DotsButton:hover {
    background-color: #f3f4f6;
}

QPushButton#DotsButton:pressed {
    background-color: #e5e7eb;
}

QMessageBox {
    background-color: #ffffff;
}
"""


# ============================================================
# KEY MANAGEMENT DIALOG
# ============================================================

class KeyManagerDialog(QDialog):

    def __init__(self, parent=None):
        super().__init__(parent)

        self.parent_window = parent

        self.setWindowTitle("Manage API Keys")
        self.setFixedSize(400, 250)

        layout = QVBoxLayout()
        layout.setContentsMargins(24, 24, 24, 24)
        layout.setSpacing(14)

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        title = QLabel("Manage API Keys")
        title.setObjectName("TitleLabel")
        title.setAlignment(Qt.AlignCenter)

        layout.addWidget(title)

        # ----------------------------------------------------
        # ADD KEY BUTTON
        # ----------------------------------------------------

        add_button = QPushButton("Add Key")
        add_button.setMinimumHeight(42)
        add_button.setCursor(Qt.PointingHandCursor)
        add_button.clicked.connect(self.add_key)

        layout.addWidget(add_button)

        # ----------------------------------------------------
        # DELETE KEY BUTTON
        # ----------------------------------------------------

        delete_button = QPushButton("Delete Selected Key")
        delete_button.setMinimumHeight(42)
        delete_button.setCursor(Qt.PointingHandCursor)
        delete_button.clicked.connect(self.delete_key)

        layout.addWidget(delete_button)

        # ----------------------------------------------------
        # CLOSE BUTTON
        # ----------------------------------------------------

        close_button = QPushButton("Close")
        close_button.setMinimumHeight(38)
        close_button.setCursor(Qt.PointingHandCursor)
        close_button.clicked.connect(self.accept)

        layout.addWidget(close_button)

        self.setLayout(layout)

    # ========================================================
    # ADD KEY
    # ========================================================

    def add_key(self):

        nickname, ok = QInputDialog.getText(
            self,
            "Add API Key",
            "Nickname:"
        )

        if not ok:
            return

        nickname = nickname.strip()

        if not nickname:

            QMessageBox.warning(
                self,
                "Invalid Nickname",
                "Nickname cannot be empty."
            )

            return

        # ----------------------------------------------------
        # CHECK EXISTING NICKNAME
        # ----------------------------------------------------

        existing_keys = self.parent_window.keys

        if nickname in existing_keys:

            QMessageBox.warning(
                self,
                "Already Exists",
                f"The nickname '{nickname}' already exists."
            )

            return

        # ----------------------------------------------------
        # API KEY
        # ----------------------------------------------------

        key, ok = QInputDialog.getText(
            self,
            "Add API Key",
            "API Key:"
        )

        if not ok:
            return

        key = key.strip()

        if not key:

            QMessageBox.warning(
                self,
                "Invalid Key",
                "API key cannot be empty."
            )

            return

        # ----------------------------------------------------
        # SAVE TO EXCEL
        # ----------------------------------------------------

        try:

            workbook = load_workbook(filename=str(KEY_FILE))

            sheet = workbook.active

            # Find next available row
            next_row = sheet.max_row + 1

            # If workbook is empty
            if (
                sheet.max_row == 1
                and sheet.cell(1, 1).value is None
                and sheet.cell(1, 2).value is None
            ):

                sheet.cell(1, 1).value = "Nickname"
                sheet.cell(1, 2).value = "Key"

                next_row = 2

            sheet.cell(next_row, 1).value = nickname
            sheet.cell(next_row, 2).value = key

            workbook.save(filename=str(KEY_FILE))

            workbook.close()

            # Refresh main application
            self.parent_window.load_keys()

            QMessageBox.information(
                self,
                "Success",
                f"'{nickname}' was added successfully."
            )

        except PermissionError:

            QMessageBox.critical(
                self,
                "File Is Open",
                "Please close key.xlsx and try again."
            )

        except Exception as e:

            QMessageBox.critical(
                self,
                "Error",
                f"Could not add the key.\n\n"
                f"{type(e).__name__}: {e}"
            )

    # ========================================================
    # DELETE KEY
    # ========================================================

    def delete_key(self):

        # ----------------------------------------------------
        # CHECK KEYS
        # ----------------------------------------------------

        if self.parent_window.key_dropdown.count() == 0:

            QMessageBox.warning(
                self,
                "No Keys",
                "There are no keys to delete."
            )

            return

        # ----------------------------------------------------
        # GET SELECTED KEY
        # ----------------------------------------------------

        selected_nickname = self.parent_window.key_dropdown.currentText()

        # ----------------------------------------------------
        # CONFIRM
        # ----------------------------------------------------

        answer = QMessageBox.question(
            self,
            "Delete API Key",
            f"Delete '{selected_nickname}'?",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No
        )

        if answer != QMessageBox.Yes:
            return

        # ----------------------------------------------------
        # DELETE FROM EXCEL
        # ----------------------------------------------------

        try:

            workbook = load_workbook(filename=str(KEY_FILE))

            sheet = workbook.active

            deleted = False

            # Start from row 2 because row 1 is header
            for row in range(2, sheet.max_row + 1):

                cell_value = sheet.cell(row=row, column=1).value

                if cell_value is None:
                    continue

                cell_value = str(cell_value).strip()

                if cell_value == selected_nickname:

                    sheet.delete_rows(row, 1)

                    deleted = True

                    break

            if deleted:
                workbook.save(filename=str(KEY_FILE))

            workbook.close()

            # Refresh main dropdown
            self.parent_window.load_keys()

            if deleted:

                QMessageBox.information(
                    self,
                    "Success",
                    f"'{selected_nickname}' was deleted."
                )

            else:

                QMessageBox.warning(
                    self,
                    "Not Found",
                    "The selected key was not found in Excel."
                )

        except PermissionError:

            QMessageBox.critical(
                self,
                "File Is Open",
                "Please close key.xlsx and try again."
            )

        except Exception as e:

            QMessageBox.critical(
                self,
                "Error",
                f"Could not delete the key.\n\n"
                f"{type(e).__name__}: {e}"
            )


# ============================================================
# MAIN WINDOW
# ============================================================

class KeySwitcher(QMainWindow):

    def __init__(self):
        super().__init__()

        self.setWindowTitle("Claude Key Switcher")
        self.setFixedSize(500, 360)

        self.keys = {}

        self.setup_ui()

        self.load_keys()

    # ========================================================
    # UI
    # ========================================================

    def setup_ui(self):

        central_widget = QWidget()

        self.setCentralWidget(central_widget)

        main_layout = QVBoxLayout()

        main_layout.setContentsMargins(32, 32, 32, 28)
        main_layout.setSpacing(16)

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        title = QLabel("Claude Key Switcher")
        title.setObjectName("TitleLabel")
        title.setAlignment(Qt.AlignCenter)

        main_layout.addWidget(title)

        # ----------------------------------------------------
        # FORM
        # ----------------------------------------------------

        form_layout = QFormLayout()

        form_layout.setSpacing(14)
        form_layout.setLabelAlignment(Qt.AlignLeft)

        # ----------------------------------------------------
        # API KEY
        # ----------------------------------------------------

        self.key_dropdown = QComboBox()
        self.key_dropdown.setMinimumHeight(38)
        self.key_dropdown.setCursor(Qt.PointingHandCursor)

        form_layout.addRow("API Key:", self.key_dropdown)

        # ----------------------------------------------------
        # MODEL
        # ----------------------------------------------------

        self.model_dropdown = QComboBox()
        self.model_dropdown.setMinimumHeight(38)
        self.model_dropdown.setCursor(Qt.PointingHandCursor)
        self.model_dropdown.addItems(MODELS)

        form_layout.addRow("Model:", self.model_dropdown)

        main_layout.addLayout(form_layout)

        # ----------------------------------------------------
        # THREE DOT BUTTON
        # ----------------------------------------------------

        dots_button = QPushButton("⋮")
        dots_button.setObjectName("DotsButton")
        dots_button.setFixedSize(36, 36)
        dots_button.setCursor(Qt.PointingHandCursor)
        dots_button.setToolTip("Manage API keys")
        dots_button.clicked.connect(self.open_key_manager)

        dots_layout = QHBoxLayout()
        dots_layout.addStretch()
        dots_layout.addWidget(dots_button)

        main_layout.addLayout(dots_layout)

        main_layout.addStretch()

        # ----------------------------------------------------
        # APPLY BUTTON
        # ----------------------------------------------------

        apply_button = QPushButton("APPLY")
        apply_button.setObjectName("ApplyButton")
        apply_button.setMinimumHeight(48)
        apply_button.setCursor(Qt.PointingHandCursor)
        apply_button.clicked.connect(self.apply_settings)

        main_layout.addWidget(apply_button)

        # ----------------------------------------------------
        # STATUS
        # ----------------------------------------------------

        self.status_label = QLabel("Ready")
        self.status_label.setObjectName("StatusLabel")
        self.status_label.setAlignment(Qt.AlignCenter)

        main_layout.addWidget(self.status_label)

        central_widget.setLayout(main_layout)

    # ========================================================
    # STATUS HELPER
    # ========================================================

    def set_status(self, text, state="neutral"):
        """Update the status label text and re-apply its style
        so the color reflects success / error / neutral state."""

        self.status_label.setText(text)
        self.status_label.setProperty("state", state)

        # Force Qt to re-evaluate the stylesheet for this widget
        self.status_label.style().unpolish(self.status_label)
        self.status_label.style().polish(self.status_label)

    # ========================================================
    # OPEN KEY MANAGER
    # ========================================================

    def open_key_manager(self):

        dialog = KeyManagerDialog(self)

        dialog.exec()

    # ========================================================
    # LOAD KEYS
    # ========================================================

    def load_keys(self):

        self.key_dropdown.clear()

        self.keys.clear()

        try:

            if not KEY_FILE.exists():

                QMessageBox.critical(
                    self,
                    "File Not Found",
                    "Excel file was not found:\n\n"
                    f"{KEY_FILE}"
                )

                self.set_status("Excel file not found", "error")

                return

            workbook = load_workbook(
                filename=str(KEY_FILE),
                read_only=True,
                data_only=True
            )

            sheet = workbook.active

            rows = sheet.iter_rows(values_only=True)

            # Skip header
            next(rows, None)

            for row in rows:

                if not row:
                    continue

                nickname = row[0]

                key = row[1] if len(row) > 1 else None

                if nickname is None:
                    continue

                if key is None:
                    continue

                nickname = str(nickname).strip()
                key = str(key).strip()

                if not nickname or not key:
                    continue

                self.keys[nickname] = key

                self.key_dropdown.addItem(nickname)

            workbook.close()

            count = self.key_dropdown.count()

            if count == 0:
                self.set_status("No API keys found", "error")
            else:
                self.set_status(f"{count} key(s) loaded", "neutral")

        except Exception as e:

            QMessageBox.critical(
                self,
                "Error Loading Keys",
                "Could not load the Excel file.\n\n"
                f"{type(e).__name__}: {e}"
            )

            self.set_status("Error loading keys", "error")

    # ========================================================
    # APPLY SETTINGS
    # ========================================================

    def apply_settings(self):

        try:

            # ------------------------------------------------
            # SETTINGS FILE
            # ------------------------------------------------

            if not SETTINGS_FILE.exists():

                QMessageBox.critical(
                    self,
                    "File Not Found",
                    "settings.json was not found:\n\n"
                    f"{SETTINGS_FILE}"
                )

                return

            # ------------------------------------------------
            # API KEY
            # ------------------------------------------------

            if self.key_dropdown.count() == 0:

                QMessageBox.warning(
                    self,
                    "No API Key",
                    "Please add an API key first."
                )

                return

            selected_nickname = self.key_dropdown.currentText()

            selected_key = self.keys.get(selected_nickname)

            if not selected_key:

                QMessageBox.warning(
                    self,
                    "Invalid Key",
                    "The selected API key could not be found."
                )

                return

            # ------------------------------------------------
            # MODEL
            # ------------------------------------------------

            selected_model = self.model_dropdown.currentText()

            # ------------------------------------------------
            # READ SETTINGS
            # ------------------------------------------------

            with open(SETTINGS_FILE, "r", encoding="utf-8") as file:
                settings = json.load(file)

            # ------------------------------------------------
            # ENV
            # ------------------------------------------------

            if "env" not in settings:
                settings["env"] = {}

            env = settings["env"]

            # ------------------------------------------------
            # UPDATE API KEY
            # ------------------------------------------------

            env["ANTHROPIC_AUTH_TOKEN"] = selected_key

            # ------------------------------------------------
            # UPDATE MODEL
            # ------------------------------------------------

            for setting in MODEL_SETTINGS:
                env[setting] = selected_model

            # ------------------------------------------------
            # SAVE
            # ------------------------------------------------

            with open(SETTINGS_FILE, "w", encoding="utf-8") as file:
                json.dump(settings, file, indent=2, ensure_ascii=False)

            # ------------------------------------------------
            # SUCCESS
            # ------------------------------------------------

            self.set_status(
                f"Applied: {selected_nickname} | {selected_model}",
                "success"
            )

            QMessageBox.information(
                self,
                "Success",
                "Settings updated successfully!\n\n"
                f"API Key: {selected_nickname}\n"
                f"Model: {selected_model}"
            )

        except json.JSONDecodeError:

            QMessageBox.critical(
                self,
                "Invalid JSON",
                "settings.json contains invalid JSON."
            )

            self.set_status("Invalid settings.json", "error")

        except PermissionError:

            QMessageBox.critical(
                self,
                "Permission Error",
                "Windows denied access to settings.json."
            )

            self.set_status("Permission denied", "error")

        except Exception as e:

            QMessageBox.critical(
                self,
                "Error",
                f"Something went wrong:\n\n"
                f"{type(e).__name__}: {e}"
            )

            self.set_status("Error applying settings", "error")


# ============================================================
# START APPLICATION
# ============================================================

def main():

    app = QApplication(sys.argv)

    app.setStyleSheet(APP_STYLESHEET)

    window = KeySwitcher()

    window.show()

    sys.exit(app.exec())


if __name__ == "__main__":
    main()