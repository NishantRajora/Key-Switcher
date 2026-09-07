import sys
import json
from pathlib import Path

from openpyxl import load_workbook, Workbook

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QApplication,
    QComboBox,
    QDialog,
    QFormLayout,
    QHBoxLayout,
    QInputDialog,
    QLabel,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QVBoxLayout,
    QWidget,
    QListWidget,
    QLineEdit,
)


# ============================================================
# FILE PATHS
# ============================================================

SETTINGS_FILE = Path(
    r"C:\Users\Nishant\.claude\settings.json"
)

KEY_FILE = Path(
    r"C:\My Space\Data\keys.xlsx"
)

PROVIDERS_FILE = Path(__file__).with_name(
    "providers.json"
)


# ============================================================
# CLAUDE CODE SETTINGS
# ============================================================

MODEL_SETTINGS = [
    "ANTHROPIC_MODEL",
    "ANTHROPIC_SMALL_FAST_MODEL",
    "ANTHROPIC_DEFAULT_SONNET_MODEL",
    "ANTHROPIC_DEFAULT_OPUS_MODEL",
    "ANTHROPIC_DEFAULT_HAIKU_MODEL",
]


# ============================================================
# DEFAULT PROVIDERS
#
# This matches your providers.json structure.
# ============================================================

DEFAULT_PROVIDERS = {
    "providers": {
        "Ollama": {
            "base_url": "https://ollama.com",
            "models": [
                "gemma4:31b",
                "gpt-oss:120b",
                "gpt-oss:20b",
                "nemotron-3-nano:30b",
                "nemotron-3-super",
                "nemotron-3-ultra"
            ]
        },

        "Hugging Face": {
            "base_url": "https://router.huggingface.co/v1",
            "models": [
                "Qwen/Qwen3.8-27B"
            ]
        }
    }
}


# ============================================================
# PROVIDER JSON FUNCTIONS
# ============================================================

def ensure_provider_file():

    if not PROVIDERS_FILE.exists():

        PROVIDERS_FILE.write_text(
            json.dumps(
                DEFAULT_PROVIDERS,
                indent=2,
                ensure_ascii=False
            ),
            encoding="utf-8"
        )


def load_providers():

    ensure_provider_file()

    try:

        data = json.loads(
            PROVIDERS_FILE.read_text(
                encoding="utf-8"
            )
        )

        if not isinstance(
            data.get("providers"),
            dict
        ):

            raise ValueError(
                "Invalid providers.json format."
            )

        return data

    except Exception as exc:

        raise RuntimeError(
            f"Could not read providers.json:\n{exc}"
        )


def save_providers(data):

    PROVIDERS_FILE.write_text(
        json.dumps(
            data,
            indent=2,
            ensure_ascii=False
        ),
        encoding="utf-8"
    )


# ============================================================
# EXCEL FUNCTIONS
# ============================================================

def ensure_key_file():

    if not KEY_FILE.exists():

        KEY_FILE.parent.mkdir(
            parents=True,
            exist_ok=True
        )

        wb = Workbook()

        ws = wb.active

        ws.title = "Keys"

        ws.append([
            "Provider",
            "Nickname",
            "API Key"
        ])

        wb.save(
            KEY_FILE
        )

        wb.close()


# ============================================================
# PROVIDER SETTINGS DIALOG
# ============================================================

class ProviderSettingsDialog(QDialog):

    def __init__(
        self,
        parent,
        providers
    ):

        super().__init__(parent)

        self.parent_window = parent
        self.providers = providers

        self.setWindowTitle(
            "Provider Settings"
        )

        self.resize(
            650,
            450
        )

        layout = QVBoxLayout(
            self
        )

        form = QFormLayout()

        # ----------------------------------------------------
        # PROVIDER
        # ----------------------------------------------------

        self.provider_dropdown = QComboBox()

        self.provider_dropdown.addItems(
            self.providers[
                "providers"
            ].keys()
        )

        self.provider_dropdown.currentTextChanged.connect(
            self.load_provider
        )

        form.addRow(
            "Provider:",
            self.provider_dropdown
        )

        # ----------------------------------------------------
        # BASE URL
        # ----------------------------------------------------

        self.url_edit = QLineEdit()

        self.url_edit.setPlaceholderText(
            "Base URL"
        )

        form.addRow(
            "Base URL:",
            self.url_edit
        )

        layout.addLayout(
            form
        )

        # ----------------------------------------------------
        # MODELS
        # ----------------------------------------------------

        layout.addWidget(
            QLabel("Models:")
        )

        self.model_list = QListWidget()

        layout.addWidget(
            self.model_list
        )

        # ----------------------------------------------------
        # BUTTONS
        # ----------------------------------------------------

        buttons = QHBoxLayout()

        add_btn = QPushButton(
            "Add Model"
        )

        delete_btn = QPushButton(
            "Delete Model"
        )

        save_btn = QPushButton(
            "SAVE"
        )

        close_btn = QPushButton(
            "CLOSE"
        )

        add_btn.clicked.connect(
            self.add_model
        )

        delete_btn.clicked.connect(
            self.delete_model
        )

        save_btn.clicked.connect(
            self.save
        )

        close_btn.clicked.connect(
            self.reject
        )

        buttons.addWidget(
            add_btn
        )

        buttons.addWidget(
            delete_btn
        )

        buttons.addStretch()

        buttons.addWidget(
            save_btn
        )

        buttons.addWidget(
            close_btn
        )

        layout.addLayout(
            buttons
        )

        # ----------------------------------------------------
        # LOAD FIRST PROVIDER
        # ----------------------------------------------------

        if self.provider_dropdown.count():

            self.load_provider(
                self.provider_dropdown.currentText()
            )


    # ========================================================
    # LOAD PROVIDER
    # ========================================================

    def load_provider(
        self,
        provider
    ):

        data = (
            self.providers[
                "providers"
            ].get(
                provider,
                {}
            )
        )

        # Load provider's base URL
        self.url_edit.setText(
            data.get(
                "base_url",
                ""
            )
        )

        # Load provider's models
        self.model_list.clear()

        self.model_list.addItems(
            data.get(
                "models",
                []
            )
        )


    # ========================================================
    # ADD MODEL
    # ========================================================

    def add_model(self):

        model, ok = QInputDialog.getText(
            self,
            "Add Model",
            "Model name:"
        )

        if not ok:
            return

        model = model.strip()

        if not model:
            return

        existing = [
            self.model_list.item(i).text()
            for i in range(
                self.model_list.count()
            )
        ]

        if model in existing:

            QMessageBox.warning(
                self,
                "Duplicate",
                "This model already exists."
            )

            return

        self.model_list.addItem(
            model
        )


    # ========================================================
    # DELETE MODEL
    # ========================================================

    def delete_model(self):

        item = self.model_list.currentItem()

        if not item:

            QMessageBox.warning(
                self,
                "Delete Model",
                "Select a model first."
            )

            return

        self.model_list.takeItem(
            self.model_list.row(item)
        )


    # ========================================================
    # SAVE PROVIDER
    # ========================================================

    def save(self):

        provider = (
            self.provider_dropdown
            .currentText()
        )

        base_url = (
            self.url_edit
            .text()
            .strip()
        )

        models = [
            self.model_list.item(i)
            .text()
            .strip()

            for i in range(
                self.model_list.count()
            )

            if self.model_list.item(i)
            .text()
            .strip()
        ]

        # Keep EXACTLY the structure:
        #
        # "Provider": {
        #     "base_url": "...",
        #     "models": [...]
        # }

        self.providers[
            "providers"
        ][provider] = {

            "base_url": base_url,

            "models": models
        }

        try:

            save_providers(
                self.providers
            )

            self.parent_window.load_data()

            QMessageBox.information(
                self,
                "Saved",
                "Provider settings updated."
            )

            self.accept()

        except Exception as exc:

            QMessageBox.critical(
                self,
                "Error",
                f"Could not save providers.json:\n{exc}"
            )


# ============================================================
# KEY MANAGER
# ============================================================

class KeyManagerDialog(QDialog):

    def __init__(
        self,
        parent
    ):

        super().__init__(parent)

        self.parent_window = parent

        self.setWindowTitle(
            "Key Manager"
        )

        self.resize(
            350,
            200
        )

        layout = QVBoxLayout(
            self
        )

        add_btn = QPushButton(
            "Add Key"
        )

        delete_btn = QPushButton(
            "Delete Selected Key"
        )

        provider_btn = QPushButton(
            "Provider Settings"
        )

        close_btn = QPushButton(
            "Close"
        )

        add_btn.clicked.connect(
            self.add_key
        )

        delete_btn.clicked.connect(
            self.delete_key
        )

        provider_btn.clicked.connect(
            self.provider_settings
        )

        close_btn.clicked.connect(
            self.accept
        )

        layout.addWidget(
            add_btn
        )

        layout.addWidget(
            delete_btn
        )

        layout.addWidget(
            provider_btn
        )

        layout.addWidget(
            close_btn
        )


    # ========================================================
    # ADD KEY
    # ========================================================

    def add_key(self):

        providers = list(
            self.parent_window
            .providers[
                "providers"
            ]
            .keys()
        )

        if not providers:

            QMessageBox.warning(
                self,
                "No Providers",
                "No providers are configured."
            )

            return

        # ----------------------------------------------------
        # SELECT PROVIDER
        # ----------------------------------------------------

        provider, ok = QInputDialog.getItem(
            self,
            "Provider",
            "Select provider:",
            providers,
            0,
            False
        )

        if not ok:
            return

        # ----------------------------------------------------
        # NICKNAME
        # ----------------------------------------------------

        nickname, ok = QInputDialog.getText(
            self,
            "Add Key",
            "Nickname:"
        )

        if not ok:
            return

        nickname = nickname.strip()

        if not nickname:

            QMessageBox.warning(
                self,
                "Invalid",
                "Nickname cannot be empty."
            )

            return

        if nickname in self.parent_window.keys:

            QMessageBox.warning(
                self,
                "Duplicate",
                "This nickname already exists."
            )

            return

        # ----------------------------------------------------
        # API KEY
        # ----------------------------------------------------

        api_key, ok = QInputDialog.getText(
            self,
            "Add Key",
            "API Key:",
            QLineEdit.EchoMode.Password
        )

        if not ok:
            return

        api_key = api_key.strip()

        if not api_key:

            QMessageBox.warning(
                self,
                "Invalid",
                "API key cannot be empty."
            )

            return

        # ----------------------------------------------------
        # SAVE TO EXCEL
        # ----------------------------------------------------

        try:

            ensure_key_file()

            wb = load_workbook(
                KEY_FILE
            )

            ws = wb.active

            # Expected:
            #
            # Provider | Nickname | API Key

            ws.append([
                provider,
                nickname,
                api_key
            ])

            wb.save(
                KEY_FILE
            )

            wb.close()

            self.parent_window.load_data()

            QMessageBox.information(
                self,
                "Added",
                "API key added successfully."
            )

        except PermissionError:

            QMessageBox.critical(
                self,
                "Excel File Locked",
                "Close keys.xlsx in Excel and try again."
            )

        except Exception as exc:

            QMessageBox.critical(
                self,
                "Error",
                f"Could not add key:\n{exc}"
            )


    # ========================================================
    # DELETE KEY
    # ========================================================

    def delete_key(self):

        nickname = (
            self.parent_window
            .key_dropdown
            .currentText()
            .strip()
        )

        if not nickname:

            QMessageBox.warning(
                self,
                "Delete Key",
                "No key is selected."
            )

            return

        answer = QMessageBox.question(
            self,
            "Confirm Delete",
            f"Delete the key '{nickname}'?",
            QMessageBox.StandardButton.Yes
            | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No
        )

        if (
            answer
            != QMessageBox.StandardButton.Yes
        ):

            return

        try:

            wb = load_workbook(
                KEY_FILE
            )

            ws = wb.active

            target_row = None

            # ------------------------------------------------
            # Find Nickname column
            # ------------------------------------------------

            headers = [
                str(
                    ws.cell(
                        1,
                        c
                    ).value or ""
                )
                .strip()
                .lower()

                for c in range(
                    1,
                    ws.max_column + 1
                )
            ]

            try:

                nickname_col = (
                    headers.index(
                        "nickname"
                    )
                    + 1
                )

            except ValueError:

                # Backward compatibility
                nickname_col = 2

            # ------------------------------------------------
            # Find matching nickname
            # ------------------------------------------------

            for row in range(
                2,
                ws.max_row + 1
            ):

                value = ws.cell(
                    row,
                    nickname_col
                ).value

                if (
                    str(
                        value or ""
                    )
                    .strip()
                    == nickname
                ):

                    target_row = row

                    break

            if target_row is None:

                wb.close()

                QMessageBox.warning(
                    self,
                    "Not Found",
                    f"Could not find '{nickname}' in keys.xlsx."
                )

                return

            # ------------------------------------------------
            # Delete row
            # ------------------------------------------------

            ws.delete_rows(
                target_row,
                1
            )

            wb.save(
                KEY_FILE
            )

            wb.close()

            self.parent_window.load_data()

            QMessageBox.information(
                self,
                "Deleted",
                "API key deleted successfully."
            )

        except PermissionError:

            QMessageBox.critical(
                self,
                "Excel File Locked",
                "Close keys.xlsx in Excel and try again."
            )

        except Exception as exc:

            QMessageBox.critical(
                self,
                "Error",
                f"Could not delete key:\n{exc}"
            )


    # ========================================================
    # PROVIDER SETTINGS
    # ========================================================

    def provider_settings(self):

        dialog = ProviderSettingsDialog(
            self.parent_window,
            self.parent_window.providers
        )

        dialog.exec()


# ============================================================
# MAIN APPLICATION
# ============================================================

class KeySwitcher(QMainWindow):

    def __init__(self):

        super().__init__()

        self.keys = {}

        self.key_providers = {}

        self.providers = {
            "providers": {}
        }

        self.setWindowTitle(
            "Key Switcher"
        )

        self.resize(
            560,
            260
        )

        self.setup_ui()

        self.load_data()


    # ========================================================
    # UI
    # ========================================================

    def setup_ui(self):

        central = QWidget()

        self.setCentralWidget(
            central
        )

        main_layout = QVBoxLayout(
            central
        )

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        title = QLabel(
            "Key Switcher"
        )

        title.setAlignment(
            Qt.AlignmentFlag.AlignCenter
        )

        title.setStyleSheet(
            "font-size: 22px; font-weight: bold;"
        )

        main_layout.addWidget(
            title
        )

        # ----------------------------------------------------
        # FORM
        # ----------------------------------------------------

        form = QFormLayout()

        # Nickname
        self.key_dropdown = QComboBox()

        self.key_dropdown.currentTextChanged.connect(
            self.on_nickname_changed
        )

        form.addRow(
            "Nickname:",
            self.key_dropdown
        )

        # Model
        self.model_dropdown = QComboBox()

        form.addRow(
            "Model:",
            self.model_dropdown
        )

        main_layout.addLayout(
            form
        )

        # ----------------------------------------------------
        # BUTTONS
        # ----------------------------------------------------

        buttons = QHBoxLayout()

        menu_btn = QPushButton(
            "⋮"
        )

        menu_btn.setFixedWidth(
            45
        )

        menu_btn.setToolTip(
            "Key and provider management"
        )

        menu_btn.clicked.connect(
            self.open_menu
        )

        apply_btn = QPushButton(
            "APPLY"
        )

        apply_btn.setMinimumHeight(
            38
        )

        apply_btn.clicked.connect(
            self.apply_settings
        )

        buttons.addStretch()

        buttons.addWidget(
            menu_btn
        )

        buttons.addWidget(
            apply_btn
        )

        main_layout.addLayout(
            buttons
        )

        # ----------------------------------------------------
        # STATUS
        # ----------------------------------------------------

        self.status = QLabel(
            "Ready"
        )

        self.status.setAlignment(
            Qt.AlignmentFlag.AlignCenter
        )

        main_layout.addWidget(
            self.status
        )


    # ========================================================
    # LOAD EVERYTHING
    # ========================================================

    def load_data(self):

        try:

            self.providers = load_providers()

            self.load_keys()

        except Exception as exc:

            QMessageBox.critical(
                self,
                "Startup Error",
                str(exc)
            )


    # ========================================================
    # LOAD EXCEL KEYS
    # ========================================================

    def load_keys(self):

        current = (
            self.key_dropdown
            .currentText()
        )

        self.keys.clear()

        self.key_providers.clear()

        self.key_dropdown.blockSignals(
            True
        )

        self.key_dropdown.clear()

        if not KEY_FILE.exists():

            self.key_dropdown.blockSignals(
                False
            )

            return

        try:

            wb = load_workbook(
                KEY_FILE,
                read_only=True,
                data_only=True
            )

            ws = wb.active

            # ------------------------------------------------
            # Read headers
            # ------------------------------------------------

            headers = [
                str(
                    ws.cell(
                        1,
                        c
                    ).value or ""
                )
                .strip()
                .lower()

                for c in range(
                    1,
                    ws.max_column + 1
                )
            ]

            # ------------------------------------------------
            # Find columns
            # ------------------------------------------------

            try:

                provider_col = (
                    headers.index(
                        "provider"
                    )
                    + 1
                )

            except ValueError:

                provider_col = None

            try:

                nickname_col = (
                    headers.index(
                        "nickname"
                    )
                    + 1
                )

            except ValueError:

                nickname_col = 2

            try:

                key_col = (
                    headers.index(
                        "api key"
                    )
                    + 1
                )

            except ValueError:

                # Old format:
                # Nickname | Key

                try:

                    key_col = (
                        headers.index(
                            "key"
                        )
                        + 1
                    )

                except ValueError:

                    key_col = 3

            # ------------------------------------------------
            # Read rows
            # ------------------------------------------------

            for row in ws.iter_rows(
                min_row=2,
                values_only=True
            ):

                # Provider
                provider = ""

                if provider_col:

                    if (
                        len(row)
                        >= provider_col
                        and row[
                            provider_col - 1
                        ]
                    ):

                        provider = str(
                            row[
                                provider_col - 1
                            ]
                        ).strip()

                # Nickname
                if (
                    len(row)
                    < nickname_col
                    or not row[
                        nickname_col - 1
                    ]
                ):

                    continue

                nickname = str(
                    row[
                        nickname_col - 1
                    ]
                ).strip()

                # API Key
                if (
                    len(row)
                    < key_col
                    or not row[
                        key_col - 1
                    ]
                ):

                    continue

                api_key = str(
                    row[
                        key_col - 1
                    ]
                ).strip()

                if nickname and api_key:

                    self.keys[
                        nickname
                    ] = api_key

                    self.key_providers[
                        nickname
                    ] = provider

            wb.close()

        except Exception as exc:

            self.key_dropdown.blockSignals(
                False
            )

            QMessageBox.critical(
                self,
                "Excel Error",
                f"Could not read keys.xlsx:\n{exc}"
            )

            return

        # ----------------------------------------------------
        # Populate Nickname dropdown
        # ----------------------------------------------------

        self.key_dropdown.addItems(
            self.keys.keys()
        )

        if (
            current
            and current in self.keys
        ):

            self.key_dropdown.setCurrentText(
                current
            )

        elif self.key_dropdown.count():

            self.key_dropdown.setCurrentIndex(
                0
            )

        self.key_dropdown.blockSignals(
            False
        )

        self.on_nickname_changed(
            self.key_dropdown.currentText()
        )


    # ========================================================
    # NICKNAME CHANGED
    #
    # This determines the provider from Excel.
    # Then it loads that provider's models from JSON.
    # ========================================================

    def on_nickname_changed(
        self,
        nickname
    ):

        self.model_dropdown.clear()

        provider = (
            self.key_providers
            .get(
                nickname,
                ""
            )
            .strip()
        )

        # ----------------------------------------------------
        # Provider not found
        # ----------------------------------------------------

        if not provider:

            self.status.setText(
                "Provider not found for selected nickname"
            )

            return

        # ----------------------------------------------------
        # Find provider in providers.json
        # ----------------------------------------------------

        provider_data = (
            self.providers[
                "providers"
            ].get(
                provider
            )
        )

        if provider_data is None:

            self.status.setText(
                f"Provider '{provider}' not found in providers.json"
            )

            return

        # ----------------------------------------------------
        # Get models
        # ----------------------------------------------------

        models = provider_data.get(
            "models",
            []
        )

        self.model_dropdown.addItems(
            models
        )

        # ----------------------------------------------------
        # Show status
        # ----------------------------------------------------

        base_url = provider_data.get(
            "base_url",
            ""
        )

        self.status.setText(
            f"{provider} • {len(models)} model(s)"
        )


    # ========================================================
    # OPEN MENU
    # ========================================================

    def open_menu(self):

        dialog = KeyManagerDialog(
            self
        )

        dialog.exec()


    # ========================================================
    # APPLY SETTINGS
    # ========================================================

    def apply_settings(self):

        # ----------------------------------------------------
        # Selected nickname
        # ----------------------------------------------------

        nickname = (
            self.key_dropdown
            .currentText()
            .strip()
        )

        # ----------------------------------------------------
        # Selected model
        # ----------------------------------------------------

        model = (
            self.model_dropdown
            .currentText()
            .strip()
        )

        # ----------------------------------------------------
        # Validate nickname
        # ----------------------------------------------------

        if not nickname:

            QMessageBox.warning(
                self,
                "Missing Key",
                "Select a nickname."
            )

            return

        # ----------------------------------------------------
        # Validate model
        # ----------------------------------------------------

        if not model:

            QMessageBox.warning(
                self,
                "Missing Model",
                "Select a model."
            )

            return

        # ----------------------------------------------------
        # Get API key
        # ----------------------------------------------------

        api_key = self.keys.get(
            nickname
        )

        if not api_key:

            QMessageBox.warning(
                self,
                "Missing Key",
                "The selected API key was not found."
            )

            return

        # ----------------------------------------------------
        # Get provider from Excel
        # ----------------------------------------------------

        provider = (
            self.key_providers
            .get(
                nickname,
                ""
            )
            .strip()
        )

        if not provider:

            QMessageBox.warning(
                self,
                "Missing Provider",
                "No provider is associated with this nickname."
            )

            return

        # ----------------------------------------------------
        # Get provider configuration from JSON
        # ----------------------------------------------------

        provider_data = (
            self.providers[
                "providers"
            ].get(
                provider
            )
        )

        if provider_data is None:

            QMessageBox.critical(
                self,
                "Provider Not Found",
                f"Provider '{provider}' was not found "
                "in providers.json."
            )

            return

        # ----------------------------------------------------
        # Get BASE URL from JSON
        # ----------------------------------------------------

        base_url = (
            provider_data
            .get(
                "base_url",
                ""
            )
            .strip()
        )

        if not base_url:

            QMessageBox.critical(
                self,
                "Missing Base URL",
                f"No base_url is configured for "
                f"provider '{provider}'."
            )

            return

        # ----------------------------------------------------
        # Check Claude Code settings
        # ----------------------------------------------------

        if not SETTINGS_FILE.exists():

            QMessageBox.critical(
                self,
                "Settings Not Found",
                "Claude Code settings file was not found:\n"
                f"{SETTINGS_FILE}"
            )

            return

        # ----------------------------------------------------
        # UPDATE SETTINGS.JSON
        # ----------------------------------------------------

        try:

            with SETTINGS_FILE.open(
                "r",
                encoding="utf-8"
            ) as file:

                settings = json.load(
                    file
                )

            if not isinstance(
                settings,
                dict
            ):

                raise ValueError(
                    "settings.json must contain a JSON object."
                )

            # ------------------------------------------------
            # Get env
            # ------------------------------------------------

            env = settings.setdefault(
                "env",
                {}
            )

            if not isinstance(
                env,
                dict
            ):

                raise ValueError(
                    "'env' in settings.json must be an object."
                )

            # ------------------------------------------------
            # API KEY
            # ------------------------------------------------

            env[
                "ANTHROPIC_AUTH_TOKEN"
            ] = api_key

            # ------------------------------------------------
            # BASE URL
            #
            # THIS IS THE IMPORTANT FIX
            # ------------------------------------------------

            env[
                "ANTHROPIC_BASE_URL"
            ] = base_url

            # ------------------------------------------------
            # MODELS
            # ------------------------------------------------

            for setting in MODEL_SETTINGS:

                env[
                    setting
                ] = model

            # ------------------------------------------------
            # SAVE SETTINGS
            # ------------------------------------------------

            with SETTINGS_FILE.open(
                "w",
                encoding="utf-8"
            ) as file:

                json.dump(
                    settings,
                    file,
                    indent=2,
                    ensure_ascii=False
                )

                file.write(
                    "\n"
                )

            # ------------------------------------------------
            # SUCCESS
            # ------------------------------------------------

            self.status.setText(
                f"Applied: {provider} • {nickname} • {model}"
            )

            QMessageBox.information(
                self,
                "Applied",
                "Claude Code settings updated.\n\n"
                f"Provider: {provider}\n"
                f"Nickname: {nickname}\n"
                f"Model: {model}\n"
                f"Base URL: {base_url}"
            )

        # ----------------------------------------------------
        # ERRORS
        # ----------------------------------------------------

        except json.JSONDecodeError:

            QMessageBox.critical(
                self,
                "Invalid JSON",
                "settings.json contains invalid JSON."
            )

        except PermissionError:

            QMessageBox.critical(
                self,
                "Permission Error",
                "Could not write settings.json."
            )

        except Exception as exc:

            QMessageBox.critical(
                self,
                "Error",
                f"Could not update Claude Code settings:\n{exc}"
            )


# ============================================================
# MAIN
# ============================================================

def main():

    app = QApplication(
        sys.argv
    )

    window = KeySwitcher()

    window.show()

    sys.exit(
        app.exec()
    )


if __name__ == "__main__":

    main()