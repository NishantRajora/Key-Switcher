
import sys
import json
from pathlib import Path

from openpyxl import load_workbook

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QApplication,
    QComboBox,
    QFormLayout,
    QLabel,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QVBoxLayout,
    QWidget,
)


# ============================================================
# FILE LOCATIONS
# ============================================================

SETTINGS_FILE = Path(
    r"C:\Users\Nishant\.claude\settings.json"
)

KEY_FILE = Path(
    r"C:\My Space\Data\key.xlsx"
)


# ============================================================
# AVAILABLE MODELS
# ============================================================

MODELS = [
    "gemma4:31b",
    "gpt-oss:120b",
    "gpt-oss:20b",
    "nemotron-3-nano:30b",
    "nemotron-3-super",
    "nemotron-3-ultra",
]


# ============================================================
# SETTINGS THAT WILL USE THE SELECTED MODEL
# ============================================================

MODEL_SETTINGS = [
    "ANTHROPIC_MODEL",
    "ANTHROPIC_SMALL_FAST_MODEL",
    "ANTHROPIC_DEFAULT_SONNET_MODEL",
    "ANTHROPIC_DEFAULT_OPUS_MODEL",
    "ANTHROPIC_DEFAULT_HAIKU_MODEL",
]


# ============================================================
# MAIN WINDOW
# ============================================================

class KeySwitcher(QMainWindow):

    def __init__(self):
        super().__init__()

        self.setWindowTitle("Claude Key Switcher")
        self.setFixedSize(500, 350)

        # Stores:
        # {
        #     "Personal": "actual-api-key",
        #     "Work": "actual-api-key"
        # }
        self.keys = {}

        self.setup_ui()
        self.load_keys()

    # ========================================================
    # USER INTERFACE
    # ========================================================

    def setup_ui(self):

        central_widget = QWidget()
        self.setCentralWidget(central_widget)

        main_layout = QVBoxLayout()

        main_layout.setContentsMargins(
            30,
            30,
            30,
            30
        )

        main_layout.setSpacing(20)

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        title = QLabel("Claude Key Switcher")

        title.setAlignment(
            Qt.AlignCenter
        )

        title.setStyleSheet("""
            QLabel {
                font-size: 24px;
                font-weight: bold;
            }
        """)

        main_layout.addWidget(title)

        # ----------------------------------------------------
        # FORM
        # ----------------------------------------------------

        form_layout = QFormLayout()

        form_layout.setSpacing(15)

        # ----------------------------------------------------
        # API KEY DROPDOWN
        # ----------------------------------------------------

        self.key_dropdown = QComboBox()

        self.key_dropdown.setMinimumHeight(40)

        form_layout.addRow(
            "API Key:",
            self.key_dropdown
        )

        # ----------------------------------------------------
        # MODEL DROPDOWN
        # ----------------------------------------------------

        self.model_dropdown = QComboBox()

        self.model_dropdown.setMinimumHeight(40)

        self.model_dropdown.addItems(
            MODELS
        )

        form_layout.addRow(
            "Model:",
            self.model_dropdown
        )

        main_layout.addLayout(
            form_layout
        )

        # ----------------------------------------------------
        # REFRESH BUTTON
        # ----------------------------------------------------

        refresh_button = QPushButton(
            "Refresh Keys"
        )

        refresh_button.setMinimumHeight(40)

        refresh_button.clicked.connect(
            self.load_keys
        )

        main_layout.addWidget(
            refresh_button
        )

        # ----------------------------------------------------
        # APPLY BUTTON
        # ----------------------------------------------------

        apply_button = QPushButton(
            "APPLY"
        )

        apply_button.setMinimumHeight(50)

        apply_button.setStyleSheet("""
            QPushButton {
                font-size: 16px;
                font-weight: bold;
            }

            QPushButton:hover {
                background-color: #dddddd;
            }
        """)

        apply_button.clicked.connect(
            self.apply_settings
        )

        main_layout.addWidget(
            apply_button
        )

        # ----------------------------------------------------
        # STATUS
        # ----------------------------------------------------

        self.status_label = QLabel(
            "Ready"
        )

        self.status_label.setAlignment(
            Qt.AlignCenter
        )

        self.status_label.setStyleSheet("""
            QLabel {
                font-size: 14px;
            }
        """)

        main_layout.addWidget(
            self.status_label
        )

        central_widget.setLayout(
            main_layout
        )

    # ========================================================
    # LOAD KEYS FROM EXCEL
    # ========================================================

    def load_keys(self):

        self.key_dropdown.clear()

        self.keys.clear()

        try:

            # ------------------------------------------------
            # CHECK EXCEL FILE
            # ------------------------------------------------

            if not KEY_FILE.exists():

                QMessageBox.critical(
                    self,
                    "File Not Found",
                    "Excel file was not found:\n\n"
                    f"{KEY_FILE}"
                )

                self.status_label.setText(
                    "Excel file not found"
                )

                return

            # ------------------------------------------------
            # OPEN EXCEL
            # ------------------------------------------------

            workbook = load_workbook(
                filename=str(KEY_FILE),
                read_only=True,
                data_only=True
            )

            # Use the active sheet
            sheet = workbook.active

            # ------------------------------------------------
            # READ ROWS
            # ------------------------------------------------

            rows = sheet.iter_rows(
                values_only=True
            )

            # Skip first row because it contains headers
            next(rows, None)

            # ------------------------------------------------
            # READ NICKNAME + KEY
            # ------------------------------------------------

            for row in rows:

                if not row:
                    continue

                # Column A = Nickname
                nickname = row[0]

                # Column B = Key
                key = (
                    row[1]
                    if len(row) > 1
                    else None
                )

                # Ignore incomplete rows
                if nickname is None:
                    continue

                if key is None:
                    continue

                nickname = str(
                    nickname
                ).strip()

                key = str(
                    key
                ).strip()

                if not nickname:
                    continue

                if not key:
                    continue

                # Store key internally
                self.keys[nickname] = key

                # Show ONLY nickname in GUI
                self.key_dropdown.addItem(
                    nickname
                )

            workbook.close()

            # ------------------------------------------------
            # RESULT
            # ------------------------------------------------

            number_of_keys = (
                self.key_dropdown.count()
            )

            if number_of_keys == 0:

                self.status_label.setText(
                    "No API keys found"
                )

                QMessageBox.warning(
                    self,
                    "No API Keys",
                    "The Excel file was found, "
                    "but no API keys were found.\n\n"
                    "Expected format:\n\n"
                    "Nickname | Key"
                )

            else:

                self.status_label.setText(
                    f"{number_of_keys} key(s) loaded"
                )

        # ----------------------------------------------------
        # EXCEL ERROR
        # ----------------------------------------------------

        except Exception as e:

            QMessageBox.critical(
                self,
                "Error Loading Keys",
                "Could not load the Excel file.\n\n"
                f"{type(e).__name__}: {e}"
            )

            self.status_label.setText(
                "Error loading keys"
            )

    # ========================================================
    # APPLY SETTINGS
    # ========================================================

    def apply_settings(self):

        try:

            # ------------------------------------------------
            # CHECK SETTINGS FILE
            # ------------------------------------------------

            if not SETTINGS_FILE.exists():

                QMessageBox.critical(
                    self,
                    "File Not Found",
                    "settings.json was not found:\n\n"
                    f"{SETTINGS_FILE}"
                )

                return

            # ------------------------------------------------
            # CHECK API KEY LIST
            # ------------------------------------------------

            if self.key_dropdown.count() == 0:

                QMessageBox.warning(
                    self,
                    "No API Key",
                    "Please add an API key to key.xlsx."
                )

                return

            # ------------------------------------------------
            # GET SELECTED NICKNAME
            # ------------------------------------------------

            selected_nickname = (
                self.key_dropdown.currentText()
            )

            # ------------------------------------------------
            # GET ACTUAL API KEY
            # ------------------------------------------------

            selected_key = self.keys.get(
                selected_nickname
            )

            if not selected_key:

                QMessageBox.warning(
                    self,
                    "Invalid Key",
                    "The selected API key could not be found."
                )

                return

            # ------------------------------------------------
            # GET SELECTED MODEL
            # ------------------------------------------------

            selected_model = (
                self.model_dropdown.currentText()
            )

            # ------------------------------------------------
            # READ SETTINGS.JSON
            # ------------------------------------------------

            with open(
                SETTINGS_FILE,
                "r",
                encoding="utf-8"
            ) as file:

                settings = json.load(
                    file
                )

            # ------------------------------------------------
            # MAKE SURE ENV EXISTS
            # ------------------------------------------------

            if "env" not in settings:

                settings["env"] = {}

            env = settings["env"]

            # ------------------------------------------------
            # UPDATE API KEY
            # ------------------------------------------------

            env[
                "ANTHROPIC_AUTH_TOKEN"
            ] = selected_key

            # ------------------------------------------------
            # UPDATE ALL MODEL SETTINGS
            # ------------------------------------------------

            for setting in MODEL_SETTINGS:

                env[setting] = selected_model

            # ------------------------------------------------
            # SAVE SETTINGS.JSON
            # ------------------------------------------------

            with open(
                SETTINGS_FILE,
                "w",
                encoding="utf-8"
            ) as file:

                json.dump(
                    settings,
                    file,
                    indent=2,
                    ensure_ascii=False
                )

            # ------------------------------------------------
            # SUCCESS
            # ------------------------------------------------

            self.status_label.setText(
                f"Applied: {selected_nickname} | "
                f"{selected_model}"
            )

            QMessageBox.information(
                self,
                "Success",
                "Settings updated successfully!\n\n"
                f"API Key: {selected_nickname}\n"
                f"Model: {selected_model}"
            )

        # ----------------------------------------------------
        # INVALID JSON
        # ----------------------------------------------------

        except json.JSONDecodeError:

            QMessageBox.critical(
                self,
                "Invalid JSON",
                "settings.json contains invalid JSON."
            )

        # ----------------------------------------------------
        # PERMISSION ERROR
        # ----------------------------------------------------

        except PermissionError:

            QMessageBox.critical(
                self,
                "Permission Error",
                "Windows denied access to settings.json.\n\n"
                "Make sure the file is not locked and "
                "you have permission to modify it."
            )

        # ----------------------------------------------------
        # OTHER ERROR
        # ----------------------------------------------------

        except Exception as e:

            QMessageBox.critical(
                self,
                "Error",
                "Something went wrong:\n\n"
                f"{type(e).__name__}: {e}"
            )


# ============================================================
# APPLICATION START
# ============================================================

def main():

    app = QApplication(
        sys.argv
    )

    window = KeySwitcher()

    window.show()

    sys.exit(
        app.exec()
    )


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    main()